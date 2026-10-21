"""⑤ 报告生成页"""
import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parents[1]))

import io
from datetime import datetime

import numpy as np
import streamlit as st
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage, Listing

from core import llm
from core.rules import (
    CUM_LIMIT,
    LEVEL_COLORS,
    level_of_value,
    stats_to_text,
    summarize,
)
from core.style import apply_theme

st.set_page_config(page_title="⑤ 报告生成", page_icon="📄", layout="wide")
apply_theme()
st.title("📄 第 5 步：报告生成（Word）")
st.caption("自动套用 Word 模板，插入各测点预警对比图（柱色 = 绿/黄/红）、AI 研判结论和处置建议，点按钮即可下载。")

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "template.docx"


def ensure_template() -> Path:
    """第一次使用时自动生成带占位符的 Word 模板。"""
    if TEMPLATE.exists():
        return TEMPLATE
    from docx import Document

    doc = Document()
    doc.add_heading("{{ title }}", 0)
    doc.add_paragraph("报告生成日期：{{ gen_date }}    数据范围：{{ range_text }}")
    doc.add_heading("一、监测数据概况", level=1)
    doc.add_paragraph("{{ summary }}")
    doc.add_heading("二、各测点预警对比图（柱色 = 绿/黄/红）", level=1)
    doc.add_paragraph("{{ chart }}")
    doc.add_heading("三、AI 智能研判结论", level=1)
    doc.add_paragraph("{{ conclusion }}")
    doc.add_heading("四、处置建议", level=1)
    doc.add_paragraph("{{ advice }}")
    doc.save(TEMPLATE)
    return TEMPLATE


def make_chart_png(metrics) -> bytes:
    """画最新一期各测点累计变化量分组柱状图（柱色 = 预警级别），供报告插入。"""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS"]
    plt.rcParams["axes.unicode_minus"] = False

    last_day = metrics["日期"].max()
    sub = metrics[metrics["日期"] == last_day].sort_values("测点编号")
    ids = sub["测点编号"].astype(str).tolist()
    x = np.arange(len(ids))
    has_settle = bool(sub["累计沉降"].notna().any())
    width = 0.38 if has_settle else 0.5

    fig, ax = plt.subplots(figsize=(9, 4.5), dpi=150)
    c1 = [LEVEL_COLORS[level_of_value(v, CUM_LIMIT)] for v in sub["累计位移"]]
    if has_settle:
        c2 = [LEVEL_COLORS[level_of_value(v, CUM_LIMIT)] for v in sub["累计沉降"]]
        ax.bar(x - width / 2, sub["累计位移"].astype(float), width,
               color=c1, label="累计位移")
        ax.bar(x + width / 2, sub["累计沉降"].astype(float), width,
               color=c2, label="累计沉降")
    else:
        ax.bar(x, sub["累计位移"].astype(float), width, color=c1, label="累计位移")
    ax.axhline(CUM_LIMIT, color="#f2c313", linestyle="--", linewidth=1.2,
               label="黄色警戒线 ±30mm")
    ax.axhline(-CUM_LIMIT, color="#f2c313", linestyle="--", linewidth=1.2)
    ax.axhline(CUM_LIMIT * 1.5, color="#e03131", linestyle="--", linewidth=1.2,
               label="红色警戒线 ±45mm")
    ax.axhline(-CUM_LIMIT * 1.5, color="#e03131", linestyle="--", linewidth=1.2)
    data = np.concatenate(
        [sub["累计位移"].to_numpy(float)]
        + ([sub["累计沉降"].to_numpy(float)] if has_settle else []))
    data = data[~np.isnan(data)]
    peak = float(np.abs(data).max()) if data.size else 0.0
    span = max(CUM_LIMIT * 1.5, peak) * 1.2
    ax.set_ylim(-span, span)
    ax.set_xticks(x)
    ax.set_xticklabels(ids)
    ax.set_xlabel("点号")
    ax.set_ylabel("累计变化量（mm）")
    ax.set_title(f"各测点累计变化量对比（{last_day:%Y-%m-%d}，按绝对值判定，柱色 = 绿/黄/红）")
    ax.grid(alpha=0.3, axis="y")
    ax.legend(fontsize=8, loc="upper left")

    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def split_ai_text(text: str):
    """把 AI 输出拆成 (研判结论, 处置建议) 两段。"""
    if "【处置建议】" in text:
        concl, advice = text.split("【处置建议】", 1)
    else:
        concl, advice = text, ""
    concl = concl.replace("【研判结论】", "").strip()
    advice = advice.strip()
    return concl, advice


if st.session_state.get("metrics") is None:
    st.warning("还没有计算结果，请先到左侧「② 趋势计算」点一下。")
    st.stop()

m = st.session_state.metrics
stats = summarize(m)
ai = st.session_state.get("ai_result")

if ai:
    st.success("已就绪：将使用第 4 步的 AI 研判结果。")
else:
    st.warning("还没做过 AI 研判，报告将使用本地兜底结论。建议先去「④ 智能研判」点一次按钮。")

project = st.text_input("工程 / 边坡名称", value="某边坡监测项目")

if st.button("📄 生成 Word 报告", type="primary"):
    try:
        with st.spinner("正在生成报告……"):
            tpl_path = ensure_template()
            tpl = DocxTemplate(str(tpl_path))
            raw_text = ai["text"] if ai else llm.local_fallback(stats)
            conclusion, advice = split_ai_text(raw_text)
            context = {
                "title": f"{project}AI 分析报告",
                "gen_date": datetime.now().strftime("%Y-%m-%d %H:%M"),
                "range_text": (
                    f"{stats['日期范围'][0]} 至 {stats['日期范围'][1]}"
                    f"（{stats['数据模式']}监测 / {stats['测点数']} 个测点"
                    f" / {stats['记录条数']} 条记录）"
                ),
                "summary": Listing(stats_to_text(stats)),
                "chart": InlineImage(tpl, io.BytesIO(make_chart_png(m)), width=Mm(165)),
                "conclusion": conclusion,
                "advice": Listing(advice),
            }
            tpl.render(context)
            out = io.BytesIO()
            tpl.save(out)
            st.session_state.report_bytes = out.getvalue()
            st.session_state.report_name = f"{project}AI分析报告.docx"
        st.success("报告生成完成！点下面的按钮下载。")
    except Exception as e:
        st.error(f"报告生成失败：{e}")

if st.session_state.get("report_bytes"):
    st.download_button(
        "⬇️ 点击下载 Word 报告",
        data=st.session_state.report_bytes,
        file_name=st.session_state.get("report_name", "边坡监测AI分析报告.docx"),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

with st.expander("ℹ️ 说明"):
    st.markdown(
        "- Word 模板 `template.docx` 会在第一次生成时自动创建在项目文件夹里，"
        "想改报告样式可以用 Word 打开它修改（注意别删掉 `{{ }}` 占位符）。\n"
        "- 每次换了数据或重新研判后，回到本页再点一次「生成 Word 报告」即可刷新内容。"
    )
